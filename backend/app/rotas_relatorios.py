"""Relatórios para prestação de contas: o que o acervo produziu num período.

GET /api/relatorios?de=AAAA-MM-DD&ate=AAAA-MM-DD   — os números (JSON) para a tela #relatorios;
GET /api/relatorios.xlsx?de=...&ate=...           — os mesmos números numa planilha, uma aba por seção.

Seções: folhas digitalizadas no período e, SEPARADAS, as migradas/importadas do site (UX-14; por mês e por fundo), folhas e projetos publicados, pedidos de material
(recebidos, atendidos e o tempo médio de entrega), problemas (abertos e resolvidos) e acessos ao site (só com o
Google Analytics 4 direto configurado; sem ele a seção volta "sem dados" e o resto do relatório continua).
Projetos de lote de teste não entram. Datas do banco são UTC ('AAAA-MM-DD HH:MM:SS'); o filtro é pelo dia.
"""
from __future__ import annotations

import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response

from . import auth
from .db import connect

router = APIRouter(prefix="/api", tags=["relatorios"])

MESES = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]
ORIGEM = {"campvision": "Digitalizadas (CAMP Vision)", "importado": "Importadas do site", "manual": "Cadastradas à mão"}
FINALIDADE = {"publicacao": "Publicação", "pesquisa": "Pesquisa", "editorial": "Editorial", "familia": "Família",
              "exposicao": "Exposição", "outro": "Outro"}
CATEGORIA = {"autoria_divergente": "Autoria divergente", "projeto_errado": "Folha de outro projeto", "duplicata": "Folha repetida",
             "orientacao": "Orientação / rotação", "espelhado": "Digitalizada pelo verso", "codigo": "Código errado",
             "metadado": "Dado errado na ficha", "credito": "Crédito errado", "arquivo_corrompido": "Arquivo corrompido", "outro": "Outro"}
SEM_GA4 = ("Sem dados de acessos: o Google Analytics ainda não está ligado ao painel. "
           "Um admin liga em Configurações (número da propriedade GA4 e chave da conta de serviço).")


def _dia(v: str | None, nome: str) -> date | None:
    if not v:
        return None
    try:
        return date.fromisoformat(v)
    except ValueError:
        raise HTTPException(400, f"Data {nome} inválida: use AAAA-MM-DD (ex.: 2026-01-31).") from None


def periodo(de: str | None, ate: str | None) -> tuple[str, str]:
    """Sem datas: o ano corrente até hoje."""
    hoje = date.today()
    d0, d1 = _dia(de, "inicial") or date(hoje.year, 1, 1), _dia(ate, "final") or hoje
    if d1 < d0:
        raise HTTPException(400, "A data final é antes da inicial.")
    if (d1 - d0).days > 366 * 10:
        raise HTTPException(400, "Período longo demais: escolha até 10 anos.")
    return d0.isoformat(), d1.isoformat()


def rotulo_mes(mes: str) -> str:
    a, m = mes.split("-")
    return f"{MESES[int(m) - 1]}/{a}"


def _br(dia: str) -> str:
    a, m, d = dia[:10].split("-")
    return f"{d}/{m}/{a}"


def _meses(de: str, ate: str) -> list[str]:
    a, m = int(de[:4]), int(de[5:7])
    fim = (int(ate[:4]), int(ate[5:7]))
    out = []
    while (a, m) <= fim:
        out.append(f"{a:04d}-{m:02d}")
        a, m = (a + 1, 1) if m == 12 else (a, m + 1)
    return out


def _dias(v) -> float | None:
    return None if v is None else round(float(v), 1)


def _folhas(con, de: str, ate: str, meses: list[str]) -> dict:
    linhas = con.execute("""
        SELECT substr(i.criado_em,1,7) AS mes, p.fundo_codigo AS fundo, f.titulo AS titulo, i.origem AS origem, COUNT(*) AS n
          FROM item i JOIN projeto p ON p.codigo=i.projeto_codigo LEFT JOIN fundo f ON f.codigo=p.fundo_codigo
         WHERE p.lote_teste=0 AND substr(i.criado_em,1,10) BETWEEN ? AND ?
         GROUP BY 1,2,3,4""", (de, ate)).fetchall()
    vazio = lambda: {k: 0 for k in ORIGEM}  # noqa: E731
    por_mes = {m: vazio() for m in meses}
    por_fundo: dict[str, dict] = {}
    cruzado: dict[tuple[str, str], int] = {}
    for r in linhas:
        o = r["origem"] if r["origem"] in ORIGEM else "manual"
        por_mes.setdefault(r["mes"], vazio())[o] += r["n"]
        pf = por_fundo.setdefault(r["fundo"], {"fundo": r["fundo"], "titulo": r["titulo"] or r["fundo"], **vazio()})
        pf[o] += r["n"]
        cruzado[(r["mes"], r["fundo"])] = cruzado.get((r["mes"], r["fundo"]), 0) + r["n"]
    pm = [{"mes": m, "rotulo": rotulo_mes(m), **v, "total": sum(v.values())} for m, v in sorted(por_mes.items())]
    pf = sorted(({**v, "total": sum(v[k] for k in ORIGEM)} for v in por_fundo.values()), key=lambda x: (-x["total"], x["fundo"]))
    total = sum(x["total"] for x in pm)
    por_origem = {k: sum(x[k] for x in pm) for k in ORIGEM}
    meses_imp = [x["rotulo"] for x in pm if x["importado"]]
    return {
        "total": total,
        # UX-14: o número principal da prestação de contas é o que foi DIGITALIZADO no período (CAMP Vision). O que veio do site
        # (migração, evento único) e o que foi enviado à mão pelo painel ficam separados, para não inflar a produção.
        "digitalizadas": por_origem["campvision"],
        "importadas": por_origem["importado"],
        "manuais": por_origem["manual"],
        "meses_importacao": meses_imp,
        "nota_importacao": nota_importacao(por_origem["importado"], meses_imp),
        "por_origem": por_origem,
        "rotulos_origem": ORIGEM,
        "por_mes": pm,
        "por_fundo": pf,
        "por_mes_fundo": [{"mes": m, "fundo": f, "total": n} for (m, f), n in sorted(cruzado.items())],
        "acervo_ao_fim": con.execute("""SELECT COUNT(*) FROM item i JOIN projeto p ON p.codigo=i.projeto_codigo
                                         WHERE p.lote_teste=0 AND substr(i.criado_em,1,10) <= ?""", (ate,)).fetchone()[0],
    }


def nota_importacao(n: int, meses: list[str]) -> str:
    if not n:
        return ""
    quando = (" (" + ", ".join(meses) + ")") if meses else ""
    return (f"Folhas que já estavam no site e entraram no painel pela migração{quando}: evento único, não é digitalização "
            "feita no período. Por isso ficam fora do número de digitalizadas.")


def _publicacoes(con, de: str, ate: str, meses: list[str]) -> dict:
    por_mes = {m: {"folhas": 0, "projetos": 0} for m in meses}
    por_fundo: dict[str, dict] = {}

    def somar(sql: str, chave: str):
        for r in con.execute(sql, (de, ate)).fetchall():
            por_mes.setdefault(r["mes"], {"folhas": 0, "projetos": 0})[chave] += r["n"]
            pf = por_fundo.setdefault(r["fundo"], {"fundo": r["fundo"], "titulo": r["titulo"] or r["fundo"], "folhas": 0, "projetos": 0})
            pf[chave] += r["n"]

    somar("""SELECT substr(i.publicado_em,1,7) AS mes, p.fundo_codigo AS fundo, f.titulo AS titulo, COUNT(*) AS n
               FROM item i JOIN projeto p ON p.codigo=i.projeto_codigo LEFT JOIN fundo f ON f.codigo=p.fundo_codigo
              WHERE p.lote_teste=0 AND substr(i.publicado_em,1,10) BETWEEN ? AND ? GROUP BY 1,2,3""", "folhas")
    somar("""SELECT substr(p.publicado_em,1,7) AS mes, p.fundo_codigo AS fundo, f.titulo AS titulo, COUNT(*) AS n
               FROM projeto p LEFT JOIN fundo f ON f.codigo=p.fundo_codigo
              WHERE p.lote_teste=0 AND p.mesclado_em IS NULL AND p.descartado_em IS NULL AND substr(p.publicado_em,1,10) BETWEEN ? AND ? GROUP BY 1,2,3""", "projetos")
    pm = [{"mes": m, "rotulo": rotulo_mes(m), **v} for m, v in sorted(por_mes.items())]
    lista = [dict(r) for r in con.execute("""
        SELECT p.codigo, p.titulo, p.fundo_codigo AS fundo, substr(p.publicado_em,1,10) AS publicado_em,
               (SELECT COUNT(*) FROM item i WHERE i.projeto_codigo=p.codigo AND i.status_site='no_ar') AS folhas
          FROM projeto p
         WHERE p.lote_teste=0 AND p.mesclado_em IS NULL AND p.descartado_em IS NULL AND substr(p.publicado_em,1,10) BETWEEN ? AND ?
         ORDER BY p.publicado_em, p.codigo""", (de, ate)).fetchall()]
    hoje = con.execute("""SELECT
        (SELECT COUNT(*) FROM item i JOIN projeto p ON p.codigo=i.projeto_codigo WHERE p.lote_teste=0 AND i.status_site='no_ar'),
        (SELECT COUNT(*) FROM projeto WHERE lote_teste=0 AND mesclado_em IS NULL AND descartado_em IS NULL AND status_site='no_ar')""").fetchone()
    return {
        "folhas": sum(x["folhas"] for x in pm),
        "projetos": sum(x["projetos"] for x in pm),
        "por_mes": pm,
        "por_fundo": sorted(por_fundo.values(), key=lambda x: (-x["folhas"], -x["projetos"], x["fundo"])),
        "projetos_lista": lista,
        "publicados_hoje": {"folhas": hoje[0], "projetos": hoje[1]},
    }


def _pedidos(con, de: str, ate: str, meses: list[str]) -> dict:
    por_mes = {m: {"recebidos": 0, "atendidos": 0} for m in meses}
    for r in con.execute("SELECT substr(recebida_em,1,7) AS mes, COUNT(*) AS n FROM solicitacao WHERE substr(recebida_em,1,10) BETWEEN ? AND ? GROUP BY 1", (de, ate)):
        por_mes.setdefault(r["mes"], {"recebidos": 0, "atendidos": 0})["recebidos"] += r["n"]
    atend = "situacao='entregue' AND entregue_em IS NOT NULL AND substr(entregue_em,1,10) BETWEEN ? AND ?"
    for r in con.execute(f"SELECT substr(entregue_em,1,7) AS mes, COUNT(*) AS n FROM solicitacao WHERE {atend} GROUP BY 1", (de, ate)):
        por_mes.setdefault(r["mes"], {"recebidos": 0, "atendidos": 0})["atendidos"] += r["n"]
    fin = {k: {"finalidade": k, "rotulo": v, "recebidos": 0, "atendidos": 0} for k, v in FINALIDADE.items()}
    for r in con.execute("SELECT finalidade, COUNT(*) AS n FROM solicitacao WHERE substr(recebida_em,1,10) BETWEEN ? AND ? GROUP BY 1", (de, ate)):
        fin.setdefault(r["finalidade"], {"finalidade": r["finalidade"], "rotulo": r["finalidade"], "recebidos": 0, "atendidos": 0})["recebidos"] = r["n"]
    for r in con.execute(f"SELECT finalidade, COUNT(*) AS n FROM solicitacao WHERE {atend} GROUP BY 1", (de, ate)):
        fin.setdefault(r["finalidade"], {"finalidade": r["finalidade"], "rotulo": r["finalidade"], "recebidos": 0, "atendidos": 0})["atendidos"] = r["n"]
    rec = con.execute("""SELECT COUNT(*) AS n,
                                SUM(situacao='recusada') AS recusados,
                                SUM(situacao NOT IN ('entregue','recusada')) AS aguardando
                           FROM solicitacao WHERE substr(recebida_em,1,10) BETWEEN ? AND ?""", (de, ate)).fetchone()
    at = con.execute(f"SELECT COUNT(*), AVG(julianday(entregue_em) - julianday(recebida_em)) FROM solicitacao WHERE {atend}", (de, ate)).fetchone()
    pm = [{"mes": m, "rotulo": rotulo_mes(m), **v} for m, v in sorted(por_mes.items())]
    return {
        "recebidos": rec["n"], "atendidos": at[0], "recusados": rec["recusados"] or 0, "aguardando": rec["aguardando"] or 0,
        "tempo_medio_dias": _dias(at[1]),
        "por_mes": pm,
        "por_finalidade": [v for v in fin.values() if v["recebidos"] or v["atendidos"]],
    }


def _problemas(con, de: str, ate: str, meses: list[str]) -> dict:
    vazio = lambda: {"abertos": 0, "resolvidos": 0}  # noqa: E731
    por_mes = {m: vazio() for m in meses}
    tipos: dict[str, dict] = {}
    resolv = "situacao IN ('corrigido','fechado') AND resolvido_em IS NOT NULL AND substr(resolvido_em,1,10) BETWEEN ? AND ?"
    for chave, sql in (("abertos", "SELECT substr(criado_em,1,7) AS mes, categoria, COUNT(*) AS n FROM erro WHERE substr(criado_em,1,10) BETWEEN ? AND ? GROUP BY 1,2"),
                       ("resolvidos", f"SELECT substr(resolvido_em,1,7) AS mes, categoria, COUNT(*) AS n FROM erro WHERE {resolv} GROUP BY 1,2")):
        for r in con.execute(sql, (de, ate)):
            por_mes.setdefault(r["mes"], vazio())[chave] += r["n"]
            t = tipos.setdefault(r["categoria"], {"categoria": r["categoria"], "rotulo": CATEGORIA.get(r["categoria"], r["categoria"]), **vazio()})
            t[chave] += r["n"]
    pm = [{"mes": m, "rotulo": rotulo_mes(m), **v} for m, v in sorted(por_mes.items())]
    res = con.execute(f"SELECT COUNT(*), AVG(julianday(resolvido_em) - julianday(criado_em)) FROM erro WHERE {resolv}", (de, ate)).fetchone()
    desc = con.execute("SELECT COUNT(*) FROM erro WHERE situacao='ignorado' AND substr(resolvido_em,1,10) BETWEEN ? AND ?", (de, ate)).fetchone()[0]
    ainda = con.execute("SELECT COUNT(*) FROM erro WHERE situacao IN ('aberto','em_correcao') AND substr(criado_em,1,10) BETWEEN ? AND ?", (de, ate)).fetchone()[0]
    return {
        "abertos": sum(x["abertos"] for x in pm), "resolvidos": res[0], "descartados": desc,
        "ainda_abertos": ainda,
        "em_aberto_hoje": con.execute("SELECT COUNT(*) FROM erro WHERE situacao IN ('aberto','em_correcao')").fetchone()[0],
        "tempo_medio_dias": _dias(res[1]),
        "por_mes": pm,
        "por_tipo": sorted(tipos.values(), key=lambda x: (-x["abertos"], -x["resolvidos"], x["rotulo"])),
    }


def _acessos(de: str, ate: str, meses: list[str]) -> dict:
    """Só pelo GA4 direto (conta de serviço). Sem ele: 'sem dados', sem quebrar e sem chamar o site."""
    from . import ga4_direto
    from .rotas_operacao import _metric_values, _top_rows
    out = {"disponivel": False, "motivo": SEM_GA4, "origem": None, "total": None, "por_mes": [], "paginas": []}
    con = connect()
    try:
        cfg = ga4_direto.configurado(con)
    except Exception as e:  # noqa: BLE001  (chave mal preenchida)
        out["motivo"] = f"Sem dados de acessos: a configuração do Google Analytics tem um problema ({str(e)[:200]})."
        return out
    finally:
        con.close()
    if not cfg:
        return out
    fim = min(ate, date.today().isoformat())
    if fim < de:
        out["motivo"] = "Sem dados de acessos: o período ainda não começou."
        return out
    try:
        g = ga4_direto.GA4(cfg)
        mets = ["activeUsers", "sessions", "screenPageViews"]
        tot = _metric_values(g.relatorio(start_date=de, end_date=fim, metrics=mets), mets)
        por = {m: {"pessoas": 0, "visitas": 0, "paginas": 0} for m in meses}
        for r in _top_rows(g.relatorio(start_date=de, end_date=fim, metrics=mets, dimensions=["yearMonth"], limit=200)):
            ym = str(r.get("yearMonth") or "")
            if len(ym) == 6:
                por[f"{ym[:4]}-{ym[4:]}"] = {"pessoas": int(r.get("activeUsers") or 0), "visitas": int(r.get("sessions") or 0), "paginas": int(r.get("screenPageViews") or 0)}
        pags = _top_rows(g.relatorio(start_date=de, end_date=fim, metrics=["screenPageViews"], dimensions=["pagePath", "pageTitle"], limit=20))
    except Exception as e:  # noqa: BLE001
        out["motivo"] = f"Sem dados de acessos: o Google Analytics não respondeu ({str(e)[:200]})."
        return out
    out.update(disponivel=True, motivo=None, origem="Google Analytics 4",
               total={"pessoas": int(tot.get("activeUsers") or 0), "visitas": int(tot.get("sessions") or 0), "paginas": int(tot.get("screenPageViews") or 0)},
               por_mes=[{"mes": m, "rotulo": rotulo_mes(m), **v} for m, v in sorted(por.items())],
               paginas=[{"caminho": r.get("pagePath") or "", "titulo": r.get("pageTitle") or r.get("pagePath") or "",
                         "visualizacoes": int(r.get("screenPageViews") or 0)} for r in pags])
    return out


def calcular(de: str, ate: str, com_acessos: bool = True) -> dict:
    meses = _meses(de, ate)
    con = connect()
    try:
        out = {
            "periodo": {"de": de, "ate": ate, "rotulo": f"{_br(de)} a {_br(ate)}",
                        "meses": [{"mes": m, "rotulo": rotulo_mes(m)} for m in meses]},
            "gerado_em": datetime.now().strftime("%d/%m/%Y %H:%M"),
            "folhas": _folhas(con, de, ate, meses),
            "publicacoes": _publicacoes(con, de, ate, meses),
            "pedidos": _pedidos(con, de, ate, meses),
            "problemas": _problemas(con, de, ate, meses),
        }
    finally:
        con.close()
    out["acessos"] = _acessos(de, ate, meses) if com_acessos else {"disponivel": False, "motivo": SEM_GA4, "por_mes": [], "paginas": []}
    return out


@router.get("/relatorios")
def relatorios(de: str | None = None, ate: str | None = None, u: dict = Depends(auth.exige("leitura"))) -> dict:
    d0, d1 = periodo(de, ate)
    return calcular(d0, d1)


# ---------------------------------------------------------------- planilha (uma aba por seção)
def planilha(r: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    negrito, titulo = Font(bold=True), Font(bold=True, size=13)
    cabec = PatternFill("solid", fgColor="E3EAF6")
    wb = Workbook()
    abas: dict = {}

    def aba(nome: str):
        ws = wb.active if not abas else wb.create_sheet()
        ws.title = nome
        abas[nome] = {"ws": ws, "linha": 1}
        ws.cell(1, 1, f"{nome} · {r['periodo']['rotulo']}").font = titulo
        abas[nome]["linha"] = 3
        return nome

    def texto(nome: str, *vals):
        a = abas[nome]
        for j, v in enumerate(vals, 1):
            a["ws"].cell(a["linha"], j, v)
        a["linha"] += 1

    def tabela(nome: str, rotulo: str, cols: list[str], linhas: list[list], total: list | None = None):
        a = abas[nome]; ws = a["ws"]
        ws.cell(a["linha"], 1, rotulo).font = negrito
        a["linha"] += 1
        for j, c in enumerate(cols, 1):
            x = ws.cell(a["linha"], j, c); x.font = negrito; x.fill = cabec
        a["linha"] += 1
        if not linhas:
            ws.cell(a["linha"], 1, "Nada no período."); a["linha"] += 1
        for ln in linhas:
            for j, v in enumerate(ln, 1):
                ws.cell(a["linha"], j, v)
            a["linha"] += 1
        if total:
            for j, v in enumerate(total, 1):
                ws.cell(a["linha"], j, v).font = negrito
            a["linha"] += 1
        a["linha"] += 1

    f, pb, pd, pr, ac = r["folhas"], r["publicacoes"], r["pedidos"], r["problemas"], r["acessos"]
    nd = lambda v: "—" if v is None else v  # noqa: E731

    s = aba("Resumo")
    texto(s, "Período", r["periodo"]["rotulo"])
    texto(s, "Gerado em", r["gerado_em"])
    abas[s]["linha"] += 1
    tabela(s, "Números do período", ["Indicador", "Valor"], [
        ["Folhas digitalizadas no período (CAMP Vision)", f["digitalizadas"]],
        ["Folhas migradas/importadas do site (evento único)", f["importadas"]],
        ["Folhas enviadas pelo painel (cadastradas à mão)", f["manuais"]],
        ["Folhas no acervo ao fim do período", f["acervo_ao_fim"]],
        ["Folhas publicadas no período", pb["folhas"]],
        ["Projetos publicados no período", pb["projetos"]],
        ["Pedidos do site recebidos", pd["recebidos"]],
        ["Pedidos do site atendidos", pd["atendidos"]],
        ["Tempo médio de entrega (dias)", nd(pd["tempo_medio_dias"])],
        ["Problemas abertos", pr["abertos"]],
        ["Problemas resolvidos", pr["resolvidos"]],
        ["Pessoas que visitaram o site", ac["total"]["pessoas"] if ac.get("disponivel") else "sem dados"],
        ["Visitas ao site", ac["total"]["visitas"] if ac.get("disponivel") else "sem dados"],
    ])
    if f["nota_importacao"]:
        texto(s, "Nota", f["nota_importacao"])

    s = aba("Folhas")
    ori = list(ORIGEM)
    texto(s, "Digitalizadas no período (CAMP Vision)", f["digitalizadas"])
    texto(s, "Migradas/importadas do site (evento único)", f["importadas"])
    texto(s, "Enviadas pelo painel (cadastradas à mão)", f["manuais"])
    if f["nota_importacao"]:
        texto(s, "Nota", f["nota_importacao"])
    abas[s]["linha"] += 1
    tabela(s, "Folhas por mês", ["Mês", *ORIGEM.values(), "Total"],
           [[m["rotulo"], *[m[k] for k in ori], m["total"]] for m in f["por_mes"]],
           ["Total", *[f["por_origem"][k] for k in ori], f["total"]])
    tabela(s, "Folhas por fundo", ["Fundo", "Nome", *ORIGEM.values(), "Total"],
           [[x["fundo"], x["titulo"], *[x[k] for k in ori], x["total"]] for x in f["por_fundo"]],
           ["Total", "", *[f["por_origem"][k] for k in ori], f["total"]])
    fundos = [x["fundo"] for x in f["por_fundo"]]
    cruz = {(x["mes"], x["fundo"]): x["total"] for x in f["por_mes_fundo"]}
    tabela(s, "Folhas por mês e por fundo", ["Mês", *fundos, "Total"],
           [[m["rotulo"], *[cruz.get((m["mes"], fu), 0) for fu in fundos], m["total"]] for m in f["por_mes"]] if fundos else [],
           ["Total", *[x["total"] for x in f["por_fundo"]], f["total"]] if fundos else None)

    s = aba("Publicações")
    tabela(s, "Publicados por mês", ["Mês", "Folhas", "Projetos"],
           [[m["rotulo"], m["folhas"], m["projetos"]] for m in pb["por_mes"]], ["Total", pb["folhas"], pb["projetos"]])
    tabela(s, "Publicados por fundo", ["Fundo", "Nome", "Folhas", "Projetos"],
           [[x["fundo"], x["titulo"], x["folhas"], x["projetos"]] for x in pb["por_fundo"]], ["Total", "", pb["folhas"], pb["projetos"]])
    tabela(s, "Projetos publicados no período", ["Projeto", "Nome", "Fundo", "Publicado em", "Folhas publicadas hoje"],
           [[x["codigo"], x["titulo"], x["fundo"], _br(x["publicado_em"]), x["folhas"]] for x in pb["projetos_lista"]])
    texto(s, "Publicados hoje (situação atual)", f"{pb['publicados_hoje']['folhas']} folhas", f"{pb['publicados_hoje']['projetos']} projetos")

    s = aba("Pedidos")
    tabela(s, "Pedidos do site por mês", ["Mês", "Recebidos", "Atendidos"],
           [[m["rotulo"], m["recebidos"], m["atendidos"]] for m in pd["por_mes"]], ["Total", pd["recebidos"], pd["atendidos"]])
    tabela(s, "Por finalidade", ["Finalidade", "Recebidos", "Atendidos"],
           [[x["rotulo"], x["recebidos"], x["atendidos"]] for x in pd["por_finalidade"]], ["Total", pd["recebidos"], pd["atendidos"]])
    texto(s, "Tempo médio de entrega (dias)", nd(pd["tempo_medio_dias"]))
    texto(s, "Recebidos no período e recusados", pd["recusados"])
    texto(s, "Recebidos no período e ainda aguardando", pd["aguardando"])

    s = aba("Problemas")
    tabela(s, "Problemas por mês", ["Mês", "Abertos", "Resolvidos"],
           [[m["rotulo"], m["abertos"], m["resolvidos"]] for m in pr["por_mes"]], ["Total", pr["abertos"], pr["resolvidos"]])
    tabela(s, "Por tipo", ["Tipo", "Abertos", "Resolvidos"],
           [[x["rotulo"], x["abertos"], x["resolvidos"]] for x in pr["por_tipo"]], ["Total", pr["abertos"], pr["resolvidos"]])
    texto(s, "Tempo médio para resolver (dias)", nd(pr["tempo_medio_dias"]))
    texto(s, "Descartados no período", pr["descartados"])
    texto(s, "Abertos no período e ainda sem solução", pr["ainda_abertos"])
    texto(s, "Em aberto hoje (todos)", pr["em_aberto_hoje"])

    s = aba("Acessos ao site")
    if ac.get("disponivel"):
        t = ac["total"]
        tabela(s, "Acessos por mês", ["Mês", "Pessoas", "Visitas", "Páginas vistas"],
               [[m["rotulo"], m["pessoas"], m["visitas"], m["paginas"]] for m in ac["por_mes"]],
               ["Período inteiro", t["pessoas"], t["visitas"], t["paginas"]])
        tabela(s, "Páginas mais vistas", ["Página", "Endereço", "Vezes vista"],
               [[p["titulo"], p["caminho"], p["visualizacoes"]] for p in ac["paginas"]])
        texto(s, "Fonte", ac.get("origem") or "Google Analytics 4")
        texto(s, "Pessoas no período inteiro não é a soma dos meses: quem volta em outro mês conta uma vez só.")
    else:
        texto(s, ac.get("motivo") or SEM_GA4)

    for ws in wb.worksheets:
        larg: dict[int, int] = {}
        for linha in ws.iter_rows(min_row=2):
            for c in linha:
                if c.value is not None:
                    larg[c.column] = max(larg.get(c.column, 0), len(str(c.value)))
        for j, w in larg.items():
            ws.column_dimensions[get_column_letter(j)].width = min(max(w + 2, 10), 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/relatorios.xlsx")
def relatorios_xlsx(de: str | None = None, ate: str | None = None, u: dict = Depends(auth.exige("leitura"))) -> Response:
    d0, d1 = periodo(de, ate)
    nome = f"relatorio_camp_{d0}_a_{d1}.xlsx"
    return Response(planilha(calcular(d0, d1)), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{nome}"', "Cache-Control": "no-store"})
